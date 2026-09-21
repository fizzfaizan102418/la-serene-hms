from __future__ import annotations

from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import json

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .auth import require_roles
from .db import DATA_DIR, get_db
from .financial_ops import ledger_reconciliation
from .finance_controls import payment_reconciliation, revenue_report, trial_balance
from .financial_authority import folio_ledger_summary, post_folio_charge_authoritative, stay_deposit_balance
from .folio_integrity import item_has_active_charge
from .models import AuditLog, BusinessDateState, Expense, FinancialTransaction, Folio, FolioItem, LedgerEntry, Payment, Reservation, Room, User
from .business_date import get_current_business_date, lock_current_business_date
from .room_charge_accrual import accrue_room_charges_for_business_date, preview_room_charges_for_business_date

router = APIRouter(prefix="/night-audit", tags=["night-audit"])
MONEY = Decimal("0.01")
FOOD_SERVICE_CHARGE_RATE = Decimal("0.10")
FOOD_CATEGORIES = {"food", "restaurant", "room_service", "beverage", "drink", "snack"}
STAFF_SERVICE_CHARGE_ACCOUNT = "Staff Service Charges Payable"
CASH_ACCOUNTS = {"Cash"}
PACK_ROOT = DATA_DIR / "daily_closing"


class ClosingConfirm(BaseModel):
    notes: str | None = None


def money(value: Decimal | int | float | str) -> Decimal:
    return Decimal(str(value)).quantize(MONEY, rounding=ROUND_HALF_UP)


def serializable(value):
    if isinstance(value, Decimal):
        return str(money(value))
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, dict):
        return {k: serializable(v) for k, v in value.items()}
    if isinstance(value, list):
        return [serializable(v) for v in value]
    return value


def get_business_date(db: Session) -> date:
    return get_current_business_date(db)


def audit(db: Session, user_id: int, action: str, business_date: date, details: dict) -> None:
    db.add(AuditLog(user_id=user_id, action=action, entity_type="night_audit", entity_id=business_date.isoformat(), details=json.dumps(serializable(details))))


def finance_snapshot(db: Session, business_date: date) -> dict:
    reconciliation = ledger_reconciliation(business_date, db, None)
    trial = trial_balance(business_date, db, None)
    payments = payment_reconciliation(business_date, db, None)
    revenue = revenue_report(business_date, db, None)
    return {
        "status": reconciliation["reconciliation"]["status"],
        "ledger_balanced": reconciliation["ledger"]["balanced"],
        "total_debits": reconciliation["ledger"]["debits"],
        "total_credits": reconciliation["ledger"]["credits"],
        "revenue_difference": reconciliation["reconciliation"]["charge_difference"],
        "cash_difference": reconciliation["reconciliation"]["cash_difference"],
        "service_charge_collected": reconciliation["reconciliation"].get("service_charge_collected", "0.00"),
        "ledger_transactions": reconciliation["ledger"]["transactions"],
        "trial_balance": {"balanced": trial["balanced"], "total_debit": trial["total_debit"], "total_credit": trial["total_credit"], "accounts": trial["accounts"]},
        "payment_reconciliation": {"received_total": payments["received_total"], "refunded_total": payments["refunded_total"], "net_total": payments["net_total"], "methods": payments["methods"]},
        "revenue_reconciliation": {"ledger_total": reconciliation["ledger"]["revenue_credits"], "operational_total": reconciliation["authority"]["folio_charges"], "difference": reconciliation["reconciliation"]["charge_difference"], "accounts": revenue["revenue"], "total": revenue["total"]},
    }


def build_summary(db: Session, business_date: date, finance: dict | None = None):
    # Financial activity is scoped by the hotel's business date, not wall-clock
    # created_at. This keeps a room charge posted on the previous night out of
    # today's revenue even when the folio remains open for a same-day checkout.
    revenue_rows = db.execute(
        select(LedgerEntry.account, LedgerEntry.direction, func.coalesce(func.sum(LedgerEntry.amount), 0))
        .join(FinancialTransaction, FinancialTransaction.id == LedgerEntry.transaction_id)
        .where(
            FinancialTransaction.business_date == business_date,
            FinancialTransaction.status == "posted",
            LedgerEntry.account.like("Revenue - %"),
        )
        .group_by(LedgerEntry.account, LedgerEntry.direction)
    ).all()
    revenue_by_account: dict[str, Decimal] = {}
    for account, direction, amount in revenue_rows:
        signed = Decimal(amount) if direction == "credit" else -Decimal(amount)
        revenue_by_account[account] = revenue_by_account.get(account, Decimal("0.00")) + signed

    room_revenue = money(revenue_by_account.get("Revenue - room", Decimal("0.00")))
    food_revenue = money(sum(
        (amount for account, amount in revenue_by_account.items() if account.removeprefix("Revenue - ").strip().lower() in FOOD_CATEGORIES),
        Decimal("0.00"),
    ))
    other_revenue = money(sum(
        (amount for account, amount in revenue_by_account.items()
         if account != "Revenue - room"
         and account.removeprefix("Revenue - ").strip().lower() not in FOOD_CATEGORIES),
        Decimal("0.00"),
    ))
    service_charge = money(food_revenue * FOOD_SERVICE_CHARGE_RATE)

    payment_rows = db.execute(
        select(Payment.method, func.coalesce(func.sum(Payment.amount), 0))
        .join(
            FinancialTransaction,
            FinancialTransaction.reference_type == "payment",
        )
        .where(
            FinancialTransaction.reference_id == Payment.id.cast(String),
            FinancialTransaction.business_date == business_date,
            FinancialTransaction.status == "posted",
        )
        .group_by(Payment.method)
    ).all()
    payment_totals: dict[str, Decimal] = {method: Decimal(amount) for method, amount in payment_rows}
    rooms = db.scalars(select(Room)).all()
    status_counts = {s: 0 for s in ("available", "reserved", "occupied", "dirty", "out_of_order")}
    for room in rooms: status_counts[room.status] = status_counts.get(room.status, 0) + 1
    arrivals = db.scalar(select(func.count(Reservation.id)).where(Reservation.check_in == business_date, Reservation.status.in_(("reserved", "checked_in")))) or 0
    departures = db.scalar(select(func.count(Reservation.id)).where(Reservation.check_out == business_date, Reservation.status == "checked_in")) or 0
    in_house = db.scalar(select(func.count(Reservation.id)).where(Reservation.status == "checked_in")) or 0
    no_shows = db.scalar(select(func.count(Reservation.id)).where(Reservation.check_in == business_date, Reservation.status == "no_show")) or 0
    expenses = db.scalar(select(func.coalesce(func.sum(Expense.amount), 0)).where(Expense.expense_date == business_date, Expense.status == "posted")) or Decimal("0.00")
    gross_revenue = money(room_revenue + food_revenue + other_revenue)
    paid_total = money(sum(payment_totals.values(), Decimal("0.00")))
    outstanding = money(sum((folio_ledger_summary(db, folio.id).balance for folio in db.scalars(select(Folio)).all()), Decimal("0.00")))

    deposit_received = db.scalar(
        select(func.coalesce(func.sum(LedgerEntry.amount), 0))
        .join(FinancialTransaction, FinancialTransaction.id == LedgerEntry.transaction_id)
        .where(
            FinancialTransaction.business_date == business_date,
            FinancialTransaction.status == "posted",
            FinancialTransaction.transaction_type == "deposit_received",
            LedgerEntry.account == "Guest Deposits",
            LedgerEntry.direction == "credit",
        )
    ) or Decimal("0.00")
    deposit_applied = db.scalar(
        select(func.coalesce(func.sum(LedgerEntry.amount), 0))
        .join(FinancialTransaction, FinancialTransaction.id == LedgerEntry.transaction_id)
        .where(
            FinancialTransaction.business_date == business_date,
            FinancialTransaction.status == "posted",
            FinancialTransaction.transaction_type == "deposit_applied",
            LedgerEntry.account == "Guest Deposits",
            LedgerEntry.direction == "debit",
        )
    ) or Decimal("0.00")
    guest_deposit_balance = db.scalar(
        select(func.coalesce(func.sum(LedgerEntry.amount), 0))
        .join(FinancialTransaction, FinancialTransaction.id == LedgerEntry.transaction_id)
        .where(
            FinancialTransaction.status == "posted",
            LedgerEntry.account == "Guest Deposits",
        )
    ) or Decimal("0.00")
    deposit_credit = db.scalar(
        select(func.coalesce(func.sum(LedgerEntry.amount), 0))
        .join(FinancialTransaction, FinancialTransaction.id == LedgerEntry.transaction_id)
        .where(
            FinancialTransaction.status == "posted",
            LedgerEntry.account == "Guest Deposits",
            LedgerEntry.direction == "credit",
        )
    ) or Decimal("0.00")
    deposit_debit = db.scalar(
        select(func.coalesce(func.sum(LedgerEntry.amount), 0))
        .join(FinancialTransaction, FinancialTransaction.id == LedgerEntry.transaction_id)
        .where(
            FinancialTransaction.status == "posted",
            LedgerEntry.account == "Guest Deposits",
            LedgerEntry.direction == "debit",
        )
    ) or Decimal("0.00")
    guest_deposit_balance = money(Decimal(deposit_credit) - Decimal(deposit_debit))

    finance = finance or finance_snapshot(db, business_date)
    state = db.get(BusinessDateState, 1)
    posting_open = not bool(state and state.last_closed_business_date and state.last_closed_business_date >= business_date)
    return {"business_date": business_date, "generated_at": datetime.utcnow(), "posting_open": posting_open, "occupancy": {"total_rooms": len(rooms), "occupied_rooms": status_counts.get("occupied", 0), "reserved_rooms": status_counts.get("reserved", 0), "available_rooms": status_counts.get("available", 0), "dirty_rooms": status_counts.get("dirty", 0), "out_of_order_rooms": status_counts.get("out_of_order", 0), "in_house_reservations": in_house}, "movement": {"arrivals": arrivals, "departures": departures, "no_shows": no_shows}, "revenue": {"room": money(room_revenue), "food": money(food_revenue), "food_service_charge": service_charge, "other": money(other_revenue), "gross": gross_revenue}, "payments": {method: money(amount) for method, amount in payment_totals.items()} | {"total": paid_total}, "outstanding": money(outstanding), "guest_deposits": {"received_today": money(deposit_received), "applied_today": money(deposit_applied), "remaining_liability": guest_deposit_balance}, "expenses": money(expenses), "net_operating": money(gross_revenue - Decimal(expenses)), "finance": finance}


def ledger_account_delta(db: Session, business_date: date, accounts: set[str], before: bool) -> Decimal:
    query = select(LedgerEntry.direction, func.coalesce(func.sum(LedgerEntry.amount), 0)).join(FinancialTransaction, FinancialTransaction.id == LedgerEntry.transaction_id).where(LedgerEntry.account.in_(accounts), FinancialTransaction.status.in_(("posted", "reversed")))
    query = query.where(FinancialTransaction.business_date < business_date) if before else query.where(FinancialTransaction.business_date == business_date)
    rows = db.execute(query.group_by(LedgerEntry.direction)).all()
    debit = sum((Decimal(amount) for direction, amount in rows if direction == "debit"), Decimal("0.00"))
    credit = sum((Decimal(amount) for direction, amount in rows if direction == "credit"), Decimal("0.00"))
    return money(debit - credit)


def prior_closing_cash(db: Session, business_date: date) -> Decimal | None:
    """Return the last closed day's physical cash for carry-forward."""
    state = db.get(BusinessDateState, 1); prior_date = state.last_closed_business_date if state else None
    if prior_date is None or prior_date >= business_date: return None
    path = PACK_ROOT / prior_date.isoformat() / "daily-closing.json"
    if not path.exists(): return None
    try:
        payload = json.loads(path.read_text(encoding="utf-8")); report = payload.get("report", {})
        projected = report.get("pre_close", {}).get("projected_close", {}).get("cash")
        if projected is not None: return money(projected)
        historical = report.get("historical_reconciliation", {}).get("closing_cash")
        if historical is not None: return money(historical)
    except (OSError, ValueError, TypeError): return None
    return None


def build_pre_close_preview(db: Session, business_date: date, summary: dict) -> dict:
    pending = preview_room_charges_for_business_date(db, business_date=business_date)
    pending_total = money(sum((row["amount"] for row in pending), Decimal("0.00")))
    opening_cash = prior_closing_cash(db, business_date)
    if opening_cash is None:
        gross_opening_cash = ledger_account_delta(db, business_date, CASH_ACCOUNTS, True)
        opening_staff_service_charge = money(max(Decimal("0.00"), -ledger_account_delta(db, business_date, {STAFF_SERVICE_CHARGE_ACCOUNT}, True)))
        opening_cash = money(gross_opening_cash - opening_staff_service_charge)
    opening_receivables = ledger_account_delta(db, business_date, {"Guest Receivables"}, True)
    today_cash = ledger_account_delta(db, business_date, CASH_ACCOUNTS, False)
    today_service_charge = money(summary["finance"].get("service_charge_collected", "0.00"))
    today_hotel_cash = money(today_cash - today_service_charge)
    today_receivables = ledger_account_delta(db, business_date, {"Guest Receivables"}, False)
    projected_cash = money(opening_cash + today_hotel_cash)
    pending_stay_ids = {row["stay_id"] for row in pending}
    anticipated_deposit_application = money(sum((stay_deposit_balance(db, stay_id) for stay_id in pending_stay_ids), Decimal("0.00")))
    current_receivables = money(max(Decimal("0.00"), opening_receivables + today_receivables))
    anticipated_deposit_application = money(min(anticipated_deposit_application, current_receivables + pending_total))
    projected_receivables = money(max(Decimal("0.00"), current_receivables + pending_total - anticipated_deposit_application))
    return {"business_date": business_date, "opening": {"cash": money(opening_cash), "guest_receivables": money(opening_receivables), "outstanding": money(max(Decimal("0.00"), opening_receivables))}, "activity": {"room_revenue": money(summary["revenue"]["room"]), "payments_received": money(summary["payments"]["total"]), "cash_received": today_hotel_cash, "gross_cash_received": money(today_cash), "staff_service_charge": today_service_charge, "expenses": money(summary["expenses"]), "ledger_transactions": summary["finance"]["ledger_transactions"], "guest_receivables_delta": money(today_receivables)}, "pending_night_audit": {"room_charges_count": len(pending), "room_charges_total": pending_total, "anticipated_deposit_application": anticipated_deposit_application, "items": pending}, "projected_close": {"room_revenue": money(summary["revenue"]["room"] + pending_total), "cash": projected_cash, "guest_receivables": projected_receivables, "outstanding": projected_receivables, "gross_revenue": money(summary["revenue"]["gross"] + pending_total)}, "controls": {"trial_balance": "balanced" if summary["finance"]["trial_balance"]["balanced"] else "review", "revenue_difference": money(summary["finance"]["revenue_difference"]), "cash_difference": money(summary["finance"]["cash_difference"])} }


def pack_dir(business_date: date) -> Path:
    path = PACK_ROOT / business_date.isoformat(); path.mkdir(parents=True, exist_ok=True); return path


def build_json(pack: Path, summary: dict, notes: str | None, closed_by: str, closed_at: datetime) -> Path:
    path = pack / "daily-closing.json"; payload = {"report": serializable(summary), "closing": {"notes": notes, "closed_by": closed_by, "closed_at": closed_at.isoformat()}}; path.write_text(json.dumps(payload, indent=2), encoding="utf-8"); return path


def build_xlsx(pack: Path, summary: dict, notes: str | None, closed_by: str, closed_at: datetime) -> Path:
    path = pack / "daily-closing.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "Daily Closing"
    ws["A1"] = "LA SERENE HOTEL"; ws["A1"].font = Font(size=16, bold=True)
    ws["A2"] = "Daily Closing / Night Audit"; ws["A2"].font = Font(size=12, bold=True)
    ws["A3"] = "Business Date"; ws["B3"] = summary["business_date"].isoformat(); ws["A4"] = "Closed By"; ws["B4"] = closed_by; ws["A5"] = "Closed At"; ws["B5"] = closed_at.isoformat()
    pre = summary.get("pre_close") or {}; sections = [
        ("Cash Position", [("Opening cash / previous closing cash", float(pre.get("opening", {}).get("cash", 0))), ("Today's net cash movement", float(pre.get("activity", {}).get("cash_received", 0))), ("Expected closing cash in hand", float(pre.get("projected_close", {}).get("cash", 0)))]),
        ("Occupancy", [("Total rooms", summary["occupancy"]["total_rooms"]), ("Occupied rooms", summary["occupancy"]["occupied_rooms"]), ("Reserved rooms", summary["occupancy"]["reserved_rooms"]), ("Available rooms", summary["occupancy"]["available_rooms"]), ("Dirty rooms", summary["occupancy"]["dirty_rooms"]), ("Out of order", summary["occupancy"]["out_of_order_rooms"]), ("In-house reservations", summary["occupancy"]["in_house_reservations"])]),
        ("Guest Movement", [("Arrivals", summary["movement"]["arrivals"]), ("Departures", summary["movement"]["departures"]), ("No-shows", summary["movement"]["no_shows"])]),
        ("Revenue", [("Room revenue", float(summary["revenue"]["room"])), ("Food revenue", float(summary["revenue"]["food"])), ("Staff service charge (10%) — excluded from hotel revenue", float(summary["revenue"]["food_service_charge"])), ("Other revenue", float(summary["revenue"]["other"])), ("Gross revenue", float(summary["revenue"]["gross"]))]),
        ("Cashier", [(k.replace("_", " ").title(), float(v)) for k, v in summary["payments"].items()]),
        ("Finance Control", [("Status", summary["finance"]["status"]), ("Ledger transactions", summary["finance"]["ledger_transactions"]), ("Ledger balanced", summary["finance"]["ledger_balanced"]), ("Revenue difference", float(summary["finance"]["revenue_difference"])), ("Cash difference", float(summary["finance"]["cash_difference"]))]),
        ("Trial Balance", [(f"{row['account']} | debit", float(row["debit"])) for row in summary["finance"]["trial_balance"]["accounts"]] + [("Total debit", float(summary["finance"]["trial_balance"]["total_debit"])), ("Total credit", float(summary["finance"]["trial_balance"]["total_credit"])), ("Balanced", summary["finance"]["trial_balance"]["balanced"])]),
        ("Payment Reconciliation", [(f"{row['method']} | received", float(row["received"])) for row in summary["finance"]["payment_reconciliation"]["methods"]] + [("Received total", float(summary["finance"]["payment_reconciliation"]["received_total"])), ("Refunded total", float(summary["finance"]["payment_reconciliation"]["refunded_total"])), ("Net total", float(summary["finance"]["payment_reconciliation"]["net_total"]))]),
        ("Revenue Reconciliation", [("Ledger revenue", float(summary["finance"]["revenue_reconciliation"]["ledger_total"])), ("Operational folio charges", float(summary["finance"]["revenue_reconciliation"]["operational_total"])), ("Difference", float(summary["finance"]["revenue_reconciliation"]["difference"])), ("Report total", float(summary["finance"]["revenue_reconciliation"]["total"]))]),
        ("Guest Deposits", [("Received today", float(summary["guest_deposits"]["received_today"])), ("Applied today", float(summary["guest_deposits"]["applied_today"])), ("Remaining deposit liability", float(summary["guest_deposits"]["remaining_liability"]))]),
        ("Operating", [("Outstanding (end-of-day)", float(summary["outstanding"])), ("Expenses (today)", float(summary["expenses"])), ("Net operating (today)", float(summary["net_operating"]))]),
    ]
    row = 7
    for title, items in sections:
        ws.cell(row, 1, title); ws.cell(row, 1).font = Font(bold=True); ws.cell(row, 1).fill = PatternFill("solid", fgColor="EDE9E0"); row += 1
        for label, value in items: ws.cell(row, 1, label); ws.cell(row, 2, value); row += 1
        row += 1
    ws.cell(row, 1, "Closing Notes"); ws.cell(row, 1).font = Font(bold=True); ws.cell(row, 2, notes or ""); ws.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 42; ws.column_dimensions["B"].width = 30; wb.save(path); return path


def build_pdf(pack: Path, summary: dict, notes: str | None, closed_by: str, closed_at: datetime) -> Path:
    path = pack / "daily-closing.pdf"; styles = getSampleStyleSheet(); styles.add(ParagraphStyle(name="Small", parent=styles["BodyText"], fontSize=8.5, leading=11))
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=15 * mm, leftMargin=15 * mm, topMargin=14 * mm, bottomMargin=14 * mm)
    story = [Paragraph("LA SERENE HOTEL", styles["Title"]), Paragraph("Daily Closing / Night Audit", styles["Heading2"]), Paragraph(f"Business Date: {summary['business_date'].isoformat()} &nbsp;&nbsp; Closed By: {closed_by} &nbsp;&nbsp; Closed At: {closed_at.isoformat()}", styles["Small"]), Spacer(1, 5 * mm)]
    def section(title, rows): story.extend([Paragraph(title, styles["Heading3"]), Table(rows, colWidths=[95 * mm, 65 * mm], style=TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#EDE9E0")), ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"), ("GRID", (0,0), (-1,-1), 0.25, colors.HexColor("#D0CCC3")), ("ALIGN", (1,1), (1,-1), "RIGHT"), ("FONTSIZE", (0,0), (-1,-1), 8.5), ("BOTTOMPADDING", (0,0), (-1,-1), 4), ("TOPPADDING", (0,0), (-1,-1), 4)])), Spacer(1, 3 * mm)])
    pre = summary.get("pre_close") or {}; opening_cash = pre.get("opening", {}).get("cash", Decimal("0.00")); net_cash_movement = pre.get("activity", {}).get("cash_received", Decimal("0.00")); expected_closing_cash = pre.get("projected_close", {}).get("cash", Decimal("0.00"))
    section("Cash Position", [["Cash control", "Amount"], ["Opening cash / previous closing cash", f"{opening_cash:.2f}"], ["Today's net cash movement", f"{net_cash_movement:.2f}"], ["Expected closing cash in hand", f"{expected_closing_cash:.2f}"]])
    section("Occupancy", [["Metric", "Value"], ["Total rooms", summary["occupancy"]["total_rooms"]], ["Occupied rooms", summary["occupancy"]["occupied_rooms"]], ["Reserved rooms", summary["occupancy"]["reserved_rooms"]], ["Available rooms", summary["occupancy"]["available_rooms"]], ["Dirty rooms", summary["occupancy"]["dirty_rooms"]], ["Out of order", summary["occupancy"]["out_of_order_rooms"]], ["In-house reservations", summary["occupancy"]["in_house_reservations"]]])
    section("Guest Movement", [["Metric", "Value"], ["Arrivals", summary["movement"]["arrivals"]], ["Departures", summary["movement"]["departures"]], ["No-shows", summary["movement"]["no_shows"]]])
    section("Revenue", [["Metric", "Amount"], ["Room revenue", f"{summary['revenue']['room']:.2f}"], ["Food revenue", f"{summary['revenue']['food']:.2f}"], ["Food service charge (10%)", f"{summary['revenue']['food_service_charge']:.2f}"], ["Other revenue", f"{summary['revenue']['other']:.2f}"], ["Gross revenue", f"{summary['revenue']['gross']:.2f}"]])
    section("Cashier Collection", [["Payment Method", "Amount"]] + [[k.replace("_", " ").title(), f"{v:.2f}"] for k, v in summary["payments"].items()])
    tb = summary["finance"]["trial_balance"]; section("Trial Balance", [["Account", "Debit"]] + [[row["account"], f"{row['debit']:.2f}"] for row in tb["accounts"]] + [["Total debit", f"{tb['total_debit']:.2f}"], ["Total credit", f"{tb['total_credit']:.2f}"], ["Balanced", tb["balanced"]]])
    pr = summary["finance"]["payment_reconciliation"]; section("Payment Reconciliation", [["Method", "Received / Refunded / Net"]] + [[row["method"], f"{row['received']:.2f} / {row['refunded']:.2f} / {row['net']:.2f}"] for row in pr["methods"]] + [["Totals", f"{pr['received_total']:.2f} / {pr['refunded_total']:.2f} / {pr['net_total']:.2f}"]])
    rr = summary["finance"]["revenue_reconciliation"]; section("Revenue Reconciliation", [["Control", "Amount"], ["Ledger revenue", f"{rr['ledger_total']:.2f}"], ["Operational folio charges", f"{rr['operational_total']:.2f}"], ["Difference", f"{rr['difference']:.2f}"], ["Revenue report total", f"{rr['total']:.2f}"]])
    gd = summary["guest_deposits"]; section("Guest Deposits", [["Control", "Amount"], ["Received today", f"{gd['received_today']:.2f}"], ["Applied today", f"{gd['applied_today']:.2f}"], ["Remaining deposit liability", f"{gd['remaining_liability']:.2f}"]])
    section("Finance Control", [["Control", "Result"], ["Overall status", summary["finance"]["status"]], ["Ledger balanced", summary["finance"]["ledger_balanced"]], ["Cash difference", f"{summary['finance']['cash_difference']:.2f}"], ["Revenue difference", f"{summary['finance']['revenue_difference']:.2f}"], ["Ledger transactions", summary["finance"]["ledger_transactions"]]])
    section("Operating", [["Metric", "Amount"], ["Outstanding (end-of-day)", f"{summary['outstanding']:.2f}"], ["Expenses (today)", f"{summary['expenses']:.2f}"], ["Net operating (today)", f"{summary['net_operating']:.2f}"]])
    story.extend([Paragraph("Closing Notes", styles["Heading3"]), Paragraph((notes or "No closing notes recorded.").replace("\n", "<br/>"), styles["BodyText"]), Spacer(1, 8 * mm), Table([["Prepared / Closed By", "Head Office Received / Verified"], [closed_by, ""], ["Signature: __________________________", "Signature: __________________________"]], colWidths=[80 * mm, 80 * mm], style=TableStyle([("GRID", (0,0), (-1,-1), 0.5, colors.HexColor("#D0CCC3")), ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"), ("TOPPADDING", (0,0), (-1,-1), 7), ("BOTTOMPADDING", (0,0), (-1,-1), 7)]))])
    doc.build(story); return path


def create_pack(summary: dict, notes: str | None, closed_by: str, closed_at: datetime) -> dict[str, str]:
    pack = pack_dir(summary["business_date"]); files = {"json": build_json(pack, summary, notes, closed_by, closed_at), "xlsx": build_xlsx(pack, summary, notes, closed_by, closed_at), "pdf": build_pdf(pack, summary, notes, closed_by, closed_at)}; return {kind: f.name for kind, f in files.items()}


def get_pack_file(business_date: date, filename: str) -> Path:
    allowed = {"daily-closing.pdf", "daily-closing.xlsx", "daily-closing.json"}
    if filename not in allowed: raise HTTPException(status_code=400, detail="Invalid closing pack file")
    path = pack_dir(business_date) / filename
    if not path.exists(): raise HTTPException(status_code=404, detail="Closing pack has not been generated for this business date")
    return path


@router.get("/preview")
def preview(db: Session = Depends(get_db), _: User = Depends(require_roles("admin", "reception"))):
    business_date = get_business_date(db); summary = build_summary(db, business_date); summary["pre_close"] = build_pre_close_preview(db, business_date, summary); return summary


@router.post("/close")
def close_day(payload: ClosingConfirm | None = None, db: Session = Depends(get_db), user: User = Depends(require_roles("admin"))):
    state = lock_current_business_date(db); business_date = state.current_business_date
    if state.last_closed_business_date and state.last_closed_business_date >= business_date: raise HTTPException(status_code=409, detail=f"Business date {business_date.isoformat()} is already closed")
    active_departures = db.scalar(select(func.count(Reservation.id)).where(Reservation.status == "checked_in", Reservation.check_out <= business_date)) or 0
    if active_departures: raise HTTPException(status_code=409, detail="Active departures must be checked out before Night Audit can close the business date")
    accrued_room_charges = accrue_room_charges_for_business_date(db, business_date=business_date, created_by=user.id)
    finance = finance_snapshot(db, business_date)
    if finance["status"] != "balanced":
        db.rollback(); raise HTTPException(status_code=409, detail=serializable({"message": "Financial reconciliation requires review before Night Audit can close", "business_date": business_date, "finance": finance}))
    summary = build_summary(db, business_date, finance); summary["night_audit"] = {"room_charges_accrued": accrued_room_charges}; summary["pre_close"] = build_pre_close_preview(db, business_date, summary); closed_at = datetime.utcnow(); pack = create_pack(summary, payload.notes if payload else None, user.username, closed_at)
    state.last_closed_at = closed_at; state.last_closed_business_date = business_date; state.current_business_date = business_date + timedelta(days=1); state.opened_at = closed_at
    audit(db, user.id, "daily_close", business_date, {"business_date": business_date, "summary": summary, "notes": payload.notes if payload else None, "pack": pack, "closed_by": user.username, "closed_at": closed_at, "next_business_date": state.current_business_date})
    db.commit(); db.refresh(state)
    return {"status": "closed", "business_date": business_date, "next_business_date": state.current_business_date, "summary": summary, "pack": pack, "closed_by": user.username, "closed_at": closed_at, "download_urls": {k: f"/api/night-audit/pack/{business_date.isoformat()}/{v}" for k, v in pack.items()}}


@router.get("/pack/{business_date}/{filename}")
def download_pack(business_date: date, filename: str, _: User = Depends(require_roles("admin", "reception"))):
    path = get_pack_file(business_date, filename); media = {"daily-closing.pdf": "application/pdf", "daily-closing.xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "daily-closing.json": "application/json"}[filename]; return FileResponse(path, media_type=media, filename=filename)
